import os
import re
import sys
import json
import time
import argparse
import hashlib
import threading
from datetime import datetime, timezone, timedelta
from concurrent.futures import ThreadPoolExecutor
from urllib.parse import urlparse, unquote

import requests
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry
from dotenv import load_dotenv
from bs4 import BeautifulSoup

try:
    from zoneinfo import ZoneInfo
except ImportError:
    ZoneInfo = None

load_dotenv()

CANVAS_URL = (os.getenv("CANVAS_URL") or "").rstrip("/")
TOKEN = os.getenv("CANVAS_TOKEN")
if not CANVAS_URL or not TOKEN:
    raise SystemExit("Missing CANVAS_URL or CANVAS_TOKEN in .env")
CANVAS_HOST = urlparse(CANVAS_URL).netloc

EXCLUDE_COURSES = {
    int(x) for x in re.split(r"[,\s]+", os.getenv("EXCLUDE_COURSES", "")) if x.isdigit()
}


def _load_tz():
    name = os.getenv("LOCAL_TZ", "America/Chicago")
    if ZoneInfo:
        try:
            return ZoneInfo(name)
        except Exception:
            pass
    return timezone.utc


LOCAL_TZ = _load_tz()

OUTPUT_FILE = "canvas_data.json"
SUMMARY_FILE = "summary.md"
FULL_FILE_TMPL = "assignment_{aid}.json"
ATTACH_DIR = "attachments"

DEFAULT_LOOKBACK_DAYS = 14
ANNOUNCEMENT_LOOKBACK_DAYS = 60
STALE_DAYS = 180
MAX_FILE_MB = 25
TEXT_IN_JSON_LIMIT = 60_000
REDO_THRESHOLD = 0.70
MAX_OUTSIDE_LINKS_IN_SUMMARY = 50

COURSE_WORKERS = 3
ITEM_WORKERS = 8

CFG = {"raw": False}
SOFT_FAIL = (401, 403, 404)


_local = threading.local()
_stats_lock = threading.Lock()
STATS = {"requests": 0, "files_downloaded": 0, "files_reused": 0, "gdocs_fetched": 0}
ERRORS = []


def _session():
    s = getattr(_local, "session", None)
    if s is None:
        s = requests.Session()
        s.headers.update({"Authorization": f"Bearer {TOKEN}", "Accept": "application/json"})
        retry = Retry(
            total=5, backoff_factor=1,
            status_forcelist=(429, 500, 502, 503, 504),
            allowed_methods=("GET",), respect_retry_after_header=True,
        )
        adapter = HTTPAdapter(max_retries=retry, pool_connections=20, pool_maxsize=20)
        s.mount("https://", adapter)
        s.mount("http://", adapter)
        _local.session = s
    return s


def _log_error(where, detail):
    with _stats_lock:
        ERRORS.append({"where": where, "detail": str(detail)})


def _bump(key):
    with _stats_lock:
        STATS[key] += 1


def _url(path):
    return path if path.startswith("http") else f"{CANVAS_URL}{path}"


def _norm(params):
    if params is None:
        return []
    if isinstance(params, dict):
        return list(params.items())
    return list(params)


def _get(url, params=None):
    # canvas rate limits with a 403, not a 429
    r = None
    for attempt in range(6):
        try:
            r = _session().get(url, params=params, timeout=60)
        except requests.RequestException as e:
            _log_error(url, e)
            return None
        _bump("requests")
        if r.status_code == 403 and "rate limit exceeded" in r.text.lower():
            time.sleep(2 ** attempt)
            continue
        try:
            if float(r.headers.get("X-Rate-Limit-Remaining", 1000)) < 100:
                time.sleep(1)
        except ValueError:
            pass
        return r
    return r


def get_one(path, params=None):
    url = _url(path)
    r = _get(url, _norm(params))
    if r is None or r.status_code in SOFT_FAIL:
        return None
    if not r.ok:
        _log_error(url, f"HTTP {r.status_code}")
        return None
    try:
        return r.json()
    except ValueError:
        return None


def get_all(path, params=None):
    url = _url(path)
    params = _norm(params) + [("per_page", 100)]
    out = []
    while url:
        r = _get(url, params)
        if r is None or r.status_code in SOFT_FAIL:
            return out
        if not r.ok:
            _log_error(url, f"HTTP {r.status_code}")
            return out
        try:
            data = r.json()
        except ValueError:
            return out
        if not isinstance(data, list):
            return data
        out.extend(data)
        params = None
        url = r.links.get("next", {}).get("url")
    return out


def as_list(x):
    return x if isinstance(x, list) else []


def _download(url, path):
    r = _session().get(url, stream=True, timeout=120, headers={"Accept": "*/*"})
    _bump("requests")
    r.raise_for_status()
    tmp = path + ".part"
    with open(tmp, "wb") as f:
        for chunk in r.iter_content(1 << 16):
            f.write(chunk)
    os.replace(tmp, path)


_memo = {}
_memo_lock = threading.Lock()


def memo(key, builder):
    with _memo_lock:
        entry = _memo.get(key)
        if entry is None:
            entry = {"lock": threading.Lock(), "done": False, "value": None}
            _memo[key] = entry
    with entry["lock"]:
        if not entry["done"]:
            try:
                entry["value"] = builder()
            except Exception as e:
                _log_error(str(key), e)
                entry["value"] = None
            entry["done"] = True
    return entry["value"]


def compact(d):
    return {k: v for k, v in d.items() if v not in (None, "", [], {})}


def _strip_noise(soup):
    for t in soup(["script", "style", "noscript"]):
        t.decompose()
    return soup


def clean_html(html):
    if not html:
        return ""
    return _strip_noise(BeautifulSoup(html, "html.parser")).get_text("\n", strip=True)


def parse_dt(s):
    if not s:
        return None
    try:
        return datetime.fromisoformat(s.replace("Z", "+00:00"))
    except ValueError:
        return None


def fmt_local(dt):
    if not dt:
        return None
    return dt.astimezone(LOCAL_TZ).strftime("%a %b %d, %I:%M %p")


def hours_until(dt, now):
    return None if dt is None else round((dt - now).total_seconds() / 3600, 1)


def is_current_term(course):
    end = parse_dt((course.get("term") or {}).get("end_at"))
    return end is None or end > datetime.now(timezone.utc)


def _segment_after(s, marker):
    if not s or marker not in s:
        return None
    seg = s.split(marker, 1)[1].split("/")[0].split("?")[0].split("#")[0]
    return seg or None


def write_json(path, obj):
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(obj, f, indent=2, ensure_ascii=False, default=str)
    os.replace(tmp, path)


def write_text(path, text):
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        f.write(text)
    os.replace(tmp, path)


def parse_html(html):
    if not html:
        return "", [], [], []
    soup = _strip_noise(BeautifulSoup(html, "html.parser"))
    file_ids, slugs, external, seen = [], [], [], set()
    for tag in soup.find_all(["a", "img", "iframe", "video", "source"]):
        href = tag.get("href") or tag.get("src") or ""
        endpoint = tag.get("data-api-endpoint") or ""
        fid = _segment_after(endpoint, "/files/") or _segment_after(href, "/files/")
        if fid and fid.isdigit():
            if int(fid) not in file_ids:
                file_ids.append(int(fid))
            continue
        host = urlparse(href).netloc
        is_canvas = href.startswith("/") or host == CANVAS_HOST
        slug = _segment_after(href, "/pages/")
        if slug and is_canvas:
            if slug not in slugs:
                slugs.append(slug)
            continue
        wrapped = re.search(r"https://docs\.google\.com/[^\s\"'&<>]+", unquote(href))
        if wrapped and "google.com" not in host:
            href, is_canvas = wrapped.group(0), False
        if href.startswith("http") and not is_canvas and href not in seen:
            seen.add(href)
            external.append({"text": tag.get_text(strip=True) or f"[{tag.name}]", "url": href})
    return soup.get_text("\n", strip=True), file_ids, slugs, external


def extract_text(path):
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".pdf":
            import pdfplumber
            parts = []
            with pdfplumber.open(path) as pdf:
                for i, page in enumerate(pdf.pages, 1):
                    t = page.extract_text() or ""
                    if t.strip():
                        parts.append(f"[page {i}]\n{t}")
            if not parts:
                return None, "PDF has no text layer (probably scanned). Upload the file itself."
            return "\n\n".join(parts), None

        if ext == ".docx":
            import docx
            d = docx.Document(path)
            parts = [p.text for p in d.paragraphs if p.text.strip()]
            for table in d.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
            return ("\n".join(parts) or None), (None if parts else "Document is empty.")

        if ext == ".pptx":
            from pptx import Presentation
            prs = Presentation(path)
            parts = []
            for i, slide in enumerate(prs.slides, 1):
                texts = [sh.text_frame.text for sh in slide.shapes
                         if sh.has_text_frame and sh.text_frame.text.strip()]
                if slide.has_notes_slide:
                    notes = slide.notes_slide.notes_text_frame.text.strip()
                    if notes:
                        texts.append(f"(notes) {notes}")
                parts.append(f"[slide {i}]\n" + "\n".join(texts))
            return "\n\n".join(parts), None

        if ext in (".xlsx", ".xlsm"):
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            parts = []
            for ws in wb.worksheets:
                parts.append(f"[sheet {ws.title}]")
                for n, row in enumerate(ws.iter_rows(values_only=True)):
                    if n >= 2000:
                        parts.append("... (truncated)")
                        break
                    if any(v is not None for v in row):
                        parts.append(" | ".join("" if v is None else str(v) for v in row))
            return "\n".join(parts), None

        if ext in (".txt", ".md", ".csv", ".tsv", ".json"):
            with open(path, encoding="utf-8", errors="replace") as f:
                return f.read(), None

        if ext in (".html", ".htm"):
            with open(path, encoding="utf-8", errors="replace") as f:
                return clean_html(f.read()), None

        if ext in (".png", ".jpg", ".jpeg", ".gif", ".webp", ".heic"):
            return None, "Image file. Upload it if the content matters."
        if ext in (".doc", ".ppt", ".xls"):
            return None, "Old Office format, can't extract. Upload the file itself."
        return None, f"No text extractor for {ext or 'this file type'}."
    except ImportError as e:
        return None, f"Missing package for extraction: {e.name}"
    except Exception as e:
        return None, f"Extraction failed: {e}"


MANIFEST_FILE = os.path.join(ATTACH_DIR, "index.json")
MANIFEST = {}
_manifest_lock = threading.Lock()


def load_manifest():
    try:
        with open(MANIFEST_FILE, encoding="utf-8") as f:
            MANIFEST.update(json.load(f))
    except (OSError, ValueError):
        pass


def save_manifest():
    if MANIFEST:
        os.makedirs(ATTACH_DIR, exist_ok=True)
        write_json(MANIFEST_FILE, MANIFEST)


def safe_name(s, limit=80):
    s = re.sub(r"[^\w\-]+", "_", s or "").strip("_")
    return re.sub(r"_+", "_", s)[:limit].rstrip("_")


def claim_path(key, label, filename):
    with _manifest_lock:
        prev = MANIFEST.get(key)
        if prev and prev.get("path"):
            return prev["path"]
        stem, ext = os.path.splitext(filename)
        stem, label = safe_name(stem), safe_name(label, 60)
        if label and stem:
            stem = f"{label}_({stem})"
        elif label:
            stem = label
        stem = stem or "file"
        taken = {v.get("path") for v in MANIFEST.values()}
        path = os.path.join(ATTACH_DIR, f"{stem}{ext.lower()}")
        n = 2
        while path in taken or os.path.exists(path):
            path = os.path.join(ATTACH_DIR, f"{stem}_{n}{ext.lower()}")
            n += 1
        MANIFEST[key] = {"path": path}
        return path


def _remember(key, **fields):
    with _manifest_lock:
        MANIFEST.setdefault(key, {}).update(fields)


def _text_for(path):
    txt_path = path + ".txt"
    if os.path.exists(txt_path):
        with open(txt_path, encoding="utf-8") as f:
            return f.read(), None, txt_path
    text, note = extract_text(path)
    if text:
        write_text(txt_path, text)
    return text, note, txt_path


def _attach_text(info, text, note, txt_path):
    if text:
        info["text"] = text[:TEXT_IN_JSON_LIMIT]
        if len(text) > TEXT_IN_JSON_LIMIT:
            info["text_truncated"] = True
            info["full_text_file"] = txt_path
    if note:
        info["note"] = note
    return compact(info)


def _build_attachment(cid, fid, label):
    meta = None
    if cid:
        meta = get_one(f"/api/v1/courses/{cid}/files/{fid}")
    meta = meta or get_one(f"/api/v1/files/{fid}")
    if not meta:
        return {"file_id": fid, "note": "No access to this file with your token."}

    name = meta.get("display_name") or meta.get("filename") or f"file_{fid}"
    size = meta.get("size") or 0
    info = {"file_id": fid, "name": name, "size_kb": round(size / 1024, 1),
            "updated_at": meta.get("updated_at")}

    if meta.get("locked_for_user") or not meta.get("url"):
        info["note"] = meta.get("lock_explanation") or "File is locked or has no download link."
        return info
    if size > MAX_FILE_MB * 1024 * 1024:
        info["note"] = f"Over {MAX_FILE_MB} MB, not downloaded."
        return info

    os.makedirs(ATTACH_DIR, exist_ok=True)
    key = f"file:{fid}"
    path = claim_path(key, label, name)
    if os.path.exists(path) and MANIFEST.get(key, {}).get("size") == size:
        _bump("files_reused")
    else:
        _download(meta["url"], path)
        _bump("files_downloaded")
        if os.path.exists(path + ".txt"):
            os.remove(path + ".txt")
        _remember(key, size=size)
    info["saved_as"] = path
    return _attach_text(info, *_text_for(path))


def get_attachment(cid, fid, label=""):
    return memo(("file", fid), lambda: _build_attachment(cid, fid, label))


_public_local = threading.local()
GDOC_PUB = re.compile(r"docs\.google\.com/(document|spreadsheets|presentation)/d/e/([\w-]+)")
GDOC_ID = re.compile(r"docs\.google\.com/(document|spreadsheets|presentation)/(?:u/\d+/)?d/([\w-]{20,})")
GOOGLE_RESTRICTED_NOTE = (
    "Couldn't open it without a Google login (probably limited to your school account). "
    "Open it, then download it or paste the text."
)


def _public_session():
    s = getattr(_public_local, "session", None)
    if s is None:
        s = requests.Session()
        s.headers["User-Agent"] = "Mozilla/5.0 (easel)"
        adapter = HTTPAdapter(max_retries=Retry(total=3, backoff_factor=1,
                                                status_forcelist=(429, 500, 502, 503, 504)))
        s.mount("https://", adapter)
        _public_local.session = s
    return s


def is_google_doc(url):
    return bool(GDOC_PUB.search(url) or GDOC_ID.search(url))


def _google_candidates(url):
    m = GDOC_PUB.search(url)
    if m:
        kind, pid = m.groups()
        base = f"https://docs.google.com/{kind}/d/e/{pid}"
        if kind == "document":
            return kind, [(f"{base}/pub", ".html")]
        if kind == "spreadsheets":
            return kind, [(f"{base}/pub?output=xlsx", ".xlsx"), (f"{base}/pubhtml", ".html")]
        return kind, [(f"{base}/pub?output=pptx", ".pptx"), (f"{base}/pub?output=pdf", ".pdf"),
                      (f"{base}/pub", ".html")]
    m = GDOC_ID.search(url)
    if not m:
        return None, []
    kind, gid = m.groups()
    base = f"https://docs.google.com/{kind}/d/{gid}"
    if kind == "document":
        return kind, [(f"{base}/export?format=docx", ".docx"), (f"{base}/export?format=txt", ".txt")]
    if kind == "spreadsheets":
        return kind, [(f"{base}/export?format=xlsx", ".xlsx")]
    return kind, [(f"{base}/export/pptx", ".pptx"), (f"{base}/export/pdf", ".pdf")]


def _download_name(r, fallback):
    cd = r.headers.get("Content-Disposition", "")
    m = re.search(r"filename\*=UTF-8''([^;]+)", cd) or re.search(r'filename="?([^";]+)"?', cd)
    return unquote(m.group(1)) if m else fallback


def _build_google_doc(url, label):
    kind, candidates = _google_candidates(url)
    if not candidates:
        return None
    info = {"url": url, "kind": kind}
    os.makedirs(ATTACH_DIR, exist_ok=True)

    for fetch_url, ext in candidates:
        try:
            r = _public_session().get(fetch_url, timeout=90)
        except requests.RequestException:
            continue
        _bump("requests")
        ctype = r.headers.get("Content-Type", "")
        if not r.ok or "accounts.google.com" in r.url or "ServiceLogin" in r.url:
            continue
        if ext != ".html" and "text/html" in ctype:
            continue

        key = f"gdoc:{fetch_url}"
        if ext == ".html":
            soup = _strip_noise(BeautifulSoup(r.text, "html.parser"))
            title = soup.title.get_text(strip=True) if soup.title else kind
            node = soup.find(id="contents") or soup.body or soup
            text = node.get_text("\n", strip=True)
            if not text.strip():
                continue
            path = claim_path(key, label, f"{title}.txt")
            write_text(path, text)
            info.update(title=title, saved_as=path)
            _bump("gdocs_fetched")
            return _attach_text(info, text, None, path)

        name = _download_name(r, f"{kind}{ext}")
        path = claim_path(key, label, name)
        with open(path, "wb") as f:
            f.write(r.content)
        if os.path.exists(path + ".txt"):
            os.remove(path + ".txt")
        info.update(title=os.path.splitext(name)[0], saved_as=path)
        if ext == ".pptx":
            info["note"] = "Slide images aren't in the text. Upload the saved file if they matter."
        _bump("gdocs_fetched")
        text, note, txt_path = _text_for(path)
        if note and info.get("note"):
            note = f"{info['note']} {note}"
        return _attach_text(info, text, note or info.get("note"), txt_path)

    info["note"] = GOOGLE_RESTRICTED_NOTE
    return info


def get_google_doc(url, label=""):
    return memo(("gdoc", url), lambda: _build_google_doc(url, label))


def html_block(cid, html, label=""):
    text, fids, slugs, external = parse_html(html)
    gdocs = [get_google_doc(l["url"], label) for l in external if is_google_doc(l["url"])]
    return compact({
        "text": text,
        "attachments": [get_attachment(cid, f, label) for f in fids],
        "google_docs": [g for g in gdocs if g],
        "external_links": external,
        "linked_page_slugs": slugs,
    })


def get_page(cid, slug):
    def build():
        p = get_one(f"/api/v1/courses/{cid}/pages/{slug}")
        if not p:
            return None
        blk = html_block(cid, p.get("body"), p.get("title"))
        blk.pop("linked_page_slugs", None)
        return compact({"title": p.get("title"), "slug": p.get("url"),
                        "updated_at": p.get("updated_at"),
                        "body_text": blk.pop("text", None), **blk})
    return memo(("page", cid, slug), build)


def get_quiz(cid, qid):
    def build():
        q = get_one(f"/api/v1/courses/{cid}/quizzes/{qid}")
        if not q:
            return None
        blk = html_block(cid, q.get("description"), q.get("title"))
        blk.pop("linked_page_slugs", None)
        out = {
            "id": q.get("id"), "title": q.get("title"), "quiz_type": q.get("quiz_type"),
            "question_count": q.get("question_count"),
            "points_possible": q.get("points_possible"),
            "time_limit_min": q.get("time_limit"),
            "allowed_attempts": q.get("allowed_attempts"),
            "one_question_at_a_time": q.get("one_question_at_a_time"),
            "cant_go_back": q.get("cant_go_back"),
            "description_text": blk.pop("text", None), **blk,
        }
        resp = get_one(f"/api/v1/courses/{cid}/quizzes/{qid}/submissions")
        subs = resp.get("quiz_submissions", []) if isinstance(resp, dict) else as_list(resp)
        out["your_attempts"] = [compact({
            "attempt": s.get("attempt"), "score": s.get("score"),
            "kept_score": s.get("kept_score"), "finished_at": s.get("finished_at"),
            "state": s.get("workflow_state"),
        }) for s in subs]
        return compact(out)
    return memo(("quiz", cid, qid), build)


def _slim_entries(entries, names):
    out = []
    for e in entries or []:
        if e.get("deleted"):
            continue
        out.append(compact({
            "author": names.get(e.get("user_id")),
            "at": e.get("created_at"),
            "text": clean_html(e.get("message")),
            "replies": _slim_entries(e.get("replies"), names),
        }))
    return out


def get_discussion(cid, tid):
    def build():
        t = get_one(f"/api/v1/courses/{cid}/discussion_topics/{tid}")
        if not t:
            return None
        blk = html_block(cid, t.get("message"), t.get("title"))
        blk.pop("linked_page_slugs", None)
        atts = blk.pop("attachments", [])
        atts += [get_attachment(cid, a["id"], t.get("title")) for a in (t.get("attachments") or [])
                 if a.get("id")]
        out = {"id": t.get("id"), "title": t.get("title"),
               "require_initial_post": t.get("require_initial_post"),
               "message_text": blk.pop("text", None), "attachments": atts, **blk}
        view = get_one(f"/api/v1/courses/{cid}/discussion_topics/{tid}/view")
        if isinstance(view, dict):
            names = {p.get("id"): p.get("display_name") for p in view.get("participants", [])}
            out["entries"] = _slim_entries(view.get("view"), names)
        elif t.get("require_initial_post"):
            out["entries_note"] = "Replies are hidden until you post first."
        return compact(out)
    return memo(("discussion", cid, tid), build)


def get_rubric_criteria(cid, rid):
    def build():
        r = get_one(f"/api/v1/courses/{cid}/rubrics/{rid}")
        return (r or {}).get("data")
    return memo(("rubric", cid, rid), build)


def slim_rubric(criteria):
    return [compact({
        "criterion": c.get("description"),
        "details": c.get("long_description"),
        "points": c.get("points"),
        "ratings": [compact({"label": r.get("description"), "points": r.get("points"),
                             "details": r.get("long_description")})
                    for r in (c.get("ratings") or [])],
    }) for c in criteria or []]


def list_my_submissions(cid):
    subs = get_all(
        f"/api/v1/courses/{cid}/students/submissions",
        [("include[]", "submission_comments"), ("include[]", "rubric_assessment")],
    )
    return {s["assignment_id"]: s for s in as_list(subs) if s.get("assignment_id")}


def fetch_submission(cid, aid):
    return get_one(
        f"/api/v1/courses/{cid}/assignments/{aid}/submissions/self",
        [("include[]", "submission_comments"), ("include[]", "rubric_assessment")],
    )


def submission_status(sub):
    if not sub:
        return "not_started"
    if sub.get("excused"):
        return "excused"
    if sub.get("missing") and not sub.get("submitted_at"):
        return "missing"
    ws = sub.get("workflow_state")
    if ws == "graded" and (sub.get("submitted_at") or sub.get("score") is not None):
        return "graded"
    if ws in ("submitted", "pending_review") and sub.get("submitted_at"):
        return "submitted_awaiting_grade"
    return "not_started"


def teacher_comments_raw(sub):
    me = (sub or {}).get("user_id")
    return [c for c in (sub or {}).get("submission_comments") or []
            if c.get("comment") and c.get("author_id") != me]


def percent(score, points):
    if score is None or not points:
        return None
    return round(score / points * 100, 1)


def slim_submission(cid, sub, criteria, points_possible, label=""):
    if not sub:
        return None
    me = sub.get("user_id")
    names = {c.get("id"): c.get("description") for c in criteria or []}
    assessment = [compact({"criterion": names.get(k, k), "points": v.get("points"),
                           "comments": v.get("comments")})
                  for k, v in (sub.get("rubric_assessment") or {}).items()]
    teacher, yours = [], []
    for c in sub.get("submission_comments") or []:
        target = yours if (me and c.get("author_id") == me) else teacher
        target.append(compact({
            "author": c.get("author_name"),
            "at": c.get("created_at"),
            "text": c.get("comment"),
            "attachments": [get_attachment(cid, a["id"], f"{label}_feedback")
                            for a in (c.get("attachments") or []) if a.get("id")],
        }))
    return compact({
        "status": submission_status(sub),
        "score": sub.get("score"), "points_possible": points_possible,
        "percent": percent(sub.get("score"), points_possible), "grade": sub.get("grade"),
        "submitted_at": sub.get("submitted_at"), "graded_at": sub.get("graded_at"),
        "late": sub.get("late") or None, "missing": sub.get("missing") or None,
        "excused": sub.get("excused") or None, "attempt": sub.get("attempt"),
        "submission_type": sub.get("submission_type"),
        "teacher_comments": teacher,
        "your_comments": yours,
        "rubric_assessment": assessment,
    })


NO_SUBMIT_TYPES = {"none", "not_graded", "on_paper"}
DONE = {"graded", "submitted_awaiting_grade", "excused"}


def is_stale(due, term_start, now):
    if term_start:
        return due < term_start - timedelta(days=7)
    return due < now - timedelta(days=STALE_DAYS)


def lock_state(a, now):
    lock = parse_dt(a.get("lock_at"))
    closed = lock is not None and lock <= now
    return closed, bool(a.get("locked_for_user")) or closed


def bucket_for(a, now, cutoff, term_start, in_module):
    if not a.get("published", True):
        return "skip"
    sub = a.get("submission") or {}
    status = submission_status(sub)
    due = parse_dt(a.get("due_at"))
    closed, locked = lock_state(a, now)

    if status in DONE:
        ref = parse_dt(sub.get("submitted_at")) or parse_dt(sub.get("graded_at")) or due
        if ref and ref >= cutoff:
            return "recent_completed"
        return "graded_older" if status == "graded" else "skip"

    stypes = set(a.get("submission_types") or [])
    if stypes and stypes <= NO_SUBMIT_TYPES:
        return "skip"

    past_due = (due is not None and due <= now) or status == "missing"
    window_start = term_start or (now - timedelta(days=STALE_DAYS))

    if past_due and locked or closed:
        ref = due or parse_dt(a.get("lock_at"))
        if ref and ref >= window_start and (a.get("points_possible") or 0):
            return "missing_locked"
        return "skip"

    if due is None:
        if not (a.get("points_possible") or 0) and not in_module:
            return "skip"
        return "undated"
    if is_stale(due, term_start, now):
        return "stale"
    if past_due:
        return "missing"
    return "open"


def resubmit_info(a, sub, now):
    if submission_status(sub) != "graded":
        return None
    _, locked = lock_state(a, now)
    stypes = set(a.get("submission_types") or [])
    if locked or not stypes or stypes <= NO_SUBMIT_TYPES:
        return None
    allowed, attempt = a.get("allowed_attempts"), (sub or {}).get("attempt") or 0
    unlimited = allowed in (None, -1)
    if not unlimited and attempt >= allowed:
        return None

    reasons = []
    pts, score = a.get("points_possible") or 0, (sub or {}).get("score")
    pct = percent(score, pts)
    if pct is not None and pct < REDO_THRESHOLD * 100:
        reasons.append(f"scored {score:g}/{pts:g} ({pct}%)")
    comments = teacher_comments_raw(sub)
    if comments:
        reasons.append("teacher left feedback")
    if not reasons:
        return None
    return compact({
        "reasons": reasons,
        "attempts_left": "unlimited" if unlimited else allowed - attempt,
        "latest_teacher_comment": comments[-1].get("comment") if comments else None,
    })


def build_assignment(cid, cname, a, sub, bucket, now, module_info, group_names):
    criteria = a.get("rubric")
    rs = a.get("rubric_settings") or {}
    if not criteria and rs.get("id"):
        criteria = get_rubric_criteria(cid, rs["id"])

    blk = html_block(cid, a.get("description"), a.get("name"))
    linked = [get_page(cid, s) for s in blk.pop("linked_page_slugs", [])]

    due = parse_dt(a.get("due_at"))
    lock = parse_dt(a.get("lock_at"))
    stypes = a.get("submission_types") or []
    closed, locked = lock_state(a, now)

    out = {
        "id": a["id"],
        "name": a.get("name"),
        "course": cname,
        "course_id": cid,
        "category": group_names.get(a.get("assignment_group_id")),
        "bucket": bucket,
        "status": submission_status(sub),
        "due_at": a.get("due_at"),
        "due_local": fmt_local(due),
        "hours_until_due": hours_until(due, now),
        "lock_at": a.get("lock_at"),
        "lock_local": fmt_local(lock),
        "unlock_at": a.get("unlock_at"),
        "closed": closed or None,
        "locked_for_you": locked or None,
        "lock_explanation": clean_html(a.get("lock_explanation")) if locked else None,
        "points_possible": a.get("points_possible"),
        "submission_types": stypes,
        "allowed_attempts": a.get("allowed_attempts"),
        "url": a.get("html_url"),
        "module": module_info,
        "description_text": blk.pop("text", None),
        "attachments": blk.pop("attachments", None),
        "google_docs": blk.pop("google_docs", None),
        "external_links": blk.pop("external_links", None),
        "linked_pages": [p for p in linked if p],
        "rubric": slim_rubric(criteria),
        "rubric_points_total": rs.get("points_possible"),
        "submission": slim_submission(cid, sub, criteria, a.get("points_possible"), a.get("name")),
        "resubmit": resubmit_info(a, sub, now),
    }
    if "external_tool" in stypes:
        out["external_tool_note"] = (
            "Content lives in an external tool (often a New Quiz or publisher site) "
            "and can't be pulled through the API. Open it and paste the content."
        )
    if a.get("quiz_id"):
        out["quiz"] = get_quiz(cid, a["quiz_id"])
    topic = a.get("discussion_topic")
    if isinstance(topic, dict) and topic.get("id"):
        out["discussion"] = get_discussion(cid, topic["id"])
    if CFG["raw"]:
        out["raw"] = a
    return compact(out)


def list_courses():
    courses = get_all("/api/v1/courses", [
        ("enrollment_state", "active"),
        ("state[]", "available"),
        ("include[]", "term"),
        ("include[]", "teachers"),
        ("include[]", "syllabus_body"),
        ("include[]", "total_scores"),
        ("include[]", "current_grading_period_scores"),
    ])
    return [c for c in as_list(courses) if is_current_term(c)]


def list_student_enrollments():
    rows = get_all("/api/v1/users/self/enrollments", [
        ("type[]", "StudentEnrollment"),
        ("state[]", "active"),
        ("include[]", "current_points"),
    ])
    return {e["course_id"]: e for e in as_list(rows) if e.get("course_id")}


def grade_record(course, enrollment):
    g = (enrollment or {}).get("grades") or {}
    ce = (course.get("enrollments") or [{}])[0]
    rec = compact({
        "current_score": g.get("current_score"),
        "current_grade": g.get("current_grade"),
        "grading_period_score": ce.get("current_period_computed_current_score"),
        "grading_period_grade": ce.get("current_period_computed_current_grade"),
    })
    if not rec:
        rec["note"] = ("Teacher has course totals hidden." if course.get("hide_final_grades")
                       else "No score posted yet.")
    return rec


def list_modules(cid):
    modules = as_list(get_all(
        f"/api/v1/courses/{cid}/modules",
        [("include[]", "items"), ("include[]", "content_details")],
    ))
    for m in modules:  # big modules don't include items inline
        if "items" not in m and m.get("items_url"):
            m["items"] = as_list(get_all(m["items_url"], [("include[]", "content_details")]))
    return modules


def module_maps(modules):
    amap, qmap, dmap = {}, {}, {}
    for m in modules:
        for it in m.get("items") or []:
            info = compact({
                "module": m.get("name"),
                "module_position": m.get("position"),
                "module_state": m.get("state"),
                "item_position": it.get("position"),
                "locked": (it.get("content_details") or {}).get("locked_for_user") or None,
                "requirement": (it.get("completion_requirement") or {}).get("type"),
                "requirement_done": (it.get("completion_requirement") or {}).get("completed"),
            })
            target = {"Assignment": amap, "Quiz": qmap, "Discussion": dmap}.get(it.get("type"))
            if target is not None and it.get("content_id"):
                target.setdefault(it["content_id"], info)
    return amap, qmap, dmap


def resolve_item(cid, it, covered, status_map, module_name=""):
    itype, content_id = it.get("type"), it.get("content_id")
    out = {
        "position": it.get("position"),
        "title": it.get("title"),
        "type": itype,
        "indent": it.get("indent") or None,
        "requirement": (it.get("completion_requirement") or {}).get("type"),
        "requirement_done": (it.get("completion_requirement") or {}).get("completed"),
        "locked": (it.get("content_details") or {}).get("locked_for_user") or None,
    }
    graded = status_map.get((itype, content_id))
    if graded:
        out.update(graded)
    if (itype, content_id) in covered:
        out["see_assignments_section"] = True
    elif itype == "Page" and it.get("page_url"):
        out["page"] = get_page(cid, it["page_url"])
    elif itype == "File" and content_id:
        out["file"] = get_attachment(cid, content_id, module_name)
    elif itype == "Discussion" and content_id:
        out["discussion"] = get_discussion(cid, content_id)
    elif itype == "Quiz" and content_id:
        out["quiz"] = get_quiz(cid, content_id)
    elif itype in ("ExternalUrl", "ExternalTool"):
        out["url"] = it.get("external_url")
    elif itype == "Assignment":
        out["assignment_id"] = content_id
    return compact(out)


def slim_announcement(cid, a):
    blk = html_block(cid, a.get("message"), a.get("title"))
    blk.pop("linked_page_slugs", None)
    atts = blk.pop("attachments", [])
    atts += [get_attachment(cid, x["id"], a.get("title")) for x in (a.get("attachments") or [])
             if x.get("id")]
    return compact({
        "title": a.get("title"),
        "posted_at": a.get("posted_at"),
        "author": (a.get("author") or {}).get("display_name"),
        "text": blk.pop("text", None),
        "attachments": atts,
        **blk,
    })


def build_course(course, enrollment, now, cutoff, pool):
    cid, cname = course["id"], course.get("name")
    t0 = time.time()

    jobs = {
        "assignments": pool.submit(get_all, f"/api/v1/courses/{cid}/assignments",
                                   [("order_by", "due_at"), ("include[]", "submission")]),
        "modules": pool.submit(list_modules, cid),
        "submissions": pool.submit(list_my_submissions, cid),
        "announcements": pool.submit(get_all, "/api/v1/announcements", [
            ("context_codes[]", f"course_{cid}"),
            ("start_date", (now - timedelta(days=ANNOUNCEMENT_LOOKBACK_DAYS)).isoformat()),
            ("end_date", now.isoformat()),
        ]),
        "pages": pool.submit(get_all, f"/api/v1/courses/{cid}/pages",
                             [("sort", "updated_at"), ("order", "desc")]),
        "groups": pool.submit(get_all, f"/api/v1/courses/{cid}/assignment_groups"),
        "front_page": pool.submit(get_one, f"/api/v1/courses/{cid}/front_page"),
    }
    d = {k: f.result() for k, f in jobs.items()}

    assignments = as_list(d["assignments"])
    modules = d["modules"] or []
    groups = as_list(d["groups"])
    group_names = {g["id"]: g.get("name") for g in groups if g.get("id")}
    amap, qmap, dmap = module_maps(modules)
    term_start = (parse_dt((course.get("term") or {}).get("start_at"))
                  or parse_dt(course.get("start_at")))

    picks = []
    for a in assignments:
        topic_id = (a.get("discussion_topic") or {}).get("id")
        minfo = amap.get(a["id"]) or qmap.get(a.get("quiz_id")) or dmap.get(topic_id)
        b = bucket_for(a, now, cutoff, term_start, bool(minfo))
        if b != "skip":
            picks.append((a, b, minfo))

    subs = d["submissions"] or {}
    if not subs and picks:
        fetched = pool.map(lambda p: fetch_submission(cid, p[0]["id"]), picks)
        subs = {p[0]["id"]: s for p, s in zip(picks, fetched) if s}

    picks = [p for p in picks if p[1] != "graded_older"
             or resubmit_info(p[0], subs.get(p[0]["id"]) or p[0].get("submission"), now)]

    status_map = {}
    for a in assignments:
        st = {"assignment_id": a["id"], "status": submission_status(a.get("submission")),
              "due_at": a.get("due_at")}
        st = compact(st)
        status_map[("Assignment", a["id"])] = st
        if a.get("quiz_id"):
            status_map[("Quiz", a["quiz_id"])] = st
        topic_id = (a.get("discussion_topic") or {}).get("id")
        if topic_id:
            status_map[("Discussion", topic_id)] = st

    built = list(pool.map(
        lambda p: build_assignment(cid, cname, p[0], subs.get(p[0]["id"]) or p[0].get("submission"),
                                   p[1], now, p[2], group_names),
        picks,
    ))

    by_bucket = {k: [] for k in ("open", "missing", "missing_locked", "undated",
                                 "recent_completed", "graded_older", "stale")}
    for a in built:
        by_bucket[a["bucket"]].append(a)
    completed = by_bucket["recent_completed"] + by_bucket["graded_older"]
    completed.sort(key=lambda x: x.get("due_at") or "", reverse=True)
    by_bucket["missing_locked"].sort(key=lambda x: x.get("due_at") or x.get("lock_at") or "",
                                     reverse=True)
    by_bucket["open"].sort(key=lambda x: x.get("due_at") or "9999")
    by_bucket["missing"].sort(key=lambda x: x.get("due_at") or "")
    by_bucket["undated"].sort(key=lambda x: ((x.get("module") or {}).get("module_position") or 9999,
                                             (x.get("module") or {}).get("item_position") or 9999))

    covered = set()
    for a, _, _ in picks:
        covered.add(("Assignment", a["id"]))
        if a.get("quiz_id"):
            covered.add(("Quiz", a["quiz_id"]))
        topic_id = (a.get("discussion_topic") or {}).get("id")
        if topic_id:
            covered.add(("Discussion", topic_id))
    flat = [(i, it) for i, m in enumerate(modules) for it in (m.get("items") or [])]
    resolved = list(pool.map(
        lambda x: resolve_item(cid, x[1], covered, status_map, modules[x[0]].get("name")), flat))
    items_by_mod = [[] for _ in modules]
    for (i, _), r in zip(flat, resolved):
        items_by_mod[i].append(r)
    slim_modules = [compact({
        "name": m.get("name"),
        "position": m.get("position"),
        "state": m.get("state"),
        "unlock_at": m.get("unlock_at"),
        "completed_at": m.get("completed_at"),
        "require_sequential_progress": m.get("require_sequential_progress") or None,
        "items": items_by_mod[i],
    }) for i, m in enumerate(modules)]

    outside = []
    for i, it in flat:
        if it.get("type") in ("ExternalUrl", "ExternalTool") and it.get("external_url"):
            outside.append(compact({"title": it.get("title"), "url": it["external_url"],
                                    "module": modules[i].get("name"), "type": it.get("type")}))

    module_slugs = {it.get("page_url") for _, it in flat if it.get("type") == "Page"}
    loose = [p for p in as_list(d["pages"]) if p.get("url") and p["url"] not in module_slugs]
    loose_pages = [p for p in pool.map(lambda p: get_page(cid, p["url"]), loose) if p]

    announcements = list(pool.map(lambda a: slim_announcement(cid, a), as_list(d["announcements"])))
    fp = d["front_page"]
    front_page = get_page(cid, fp["url"]) if isinstance(fp, dict) and fp.get("url") else None
    syllabus = html_block(cid, course.get("syllabus_body"), f"{cname}_Syllabus")
    syllabus.pop("linked_page_slugs", None)

    print(f"  [{cid}] {cname}: {len(by_bucket['open'])} open, {len(by_bucket['missing'])} overdue, "
          f"{len(by_bucket['missing_locked'])} closed, {len(by_bucket['undated'])} undated, "
          f"{len(completed)} completed, {len(by_bucket['stale'])} old-dated, "
          f"{len(outside)} outside links ({time.time() - t0:.1f}s)")

    return compact({
        "id": cid,
        "name": cname,
        "term": (course.get("term") or {}).get("name"),
        "teachers": [t.get("display_name") for t in (course.get("teachers") or [])],
        "grade": grade_record(course, enrollment),
        "assignment_groups": [compact({"name": g.get("name"), "weight": g.get("group_weight")})
                              for g in groups],
        "assignments_overdue": by_bucket["missing"],
        "assignments_missing_closed": by_bucket["missing_locked"],
        "assignments_open": by_bucket["open"],
        "assignments_undated": by_bucket["undated"],
        "assignments_completed": completed,
        "assignments_old_dates": by_bucket["stale"],
        "outside_canvas_links": outside,
        "modules": slim_modules,
        "pages_not_in_modules": loose_pages,
        "announcements": announcements,
        "front_page": front_page,
        "syllabus": syllabus,
    })


STATUS_TEXT = {
    "not_started": "not started",
    "submitted_awaiting_grade": "submitted, waiting on grade",
    "graded": "graded",
    "missing": "marked missing",
    "excused": "excused",
}


def score_text(a):
    sub = a.get("submission") or {}
    if sub.get("score") is None:
        return "no score"
    pts = sub.get("points_possible")
    if not pts:
        return f"scored {sub['score']:g}"
    return f"scored {sub['score']:g}/{pts:g} ({sub.get('percent')}%)"


def _line(a, extra=""):
    pts = a.get("points_possible")
    pts_s = f"{pts:g} pts" if isinstance(pts, (int, float)) else "? pts"
    due = a.get("due_local") or "no due date"
    return (f"- **{a['course']}** | {a['name']} | {due} | {pts_s} | "
            f"{STATUS_TEXT.get(a.get('status'), a.get('status'))} | id {a['id']}{extra}")


def build_summary(dump, now):
    courses = dump["courses"]
    week_end = now + timedelta(days=7)
    L = [f"# Canvas summary, {fmt_local(now)}", ""]

    def section(title, rows):
        L.append(f"## {title}")
        L.extend(rows or ["- none"])
        L.append("")

    section("Grades", [
        f"- {c['name']}: " + (
            f"{c['grade']['current_score']}%" + (f" ({c['grade']['current_grade']})"
                                                 if c['grade'].get('current_grade') else "")
            if c["grade"].get("current_score") is not None else c["grade"].get("note", "n/a"))
        for c in courses
    ])

    overdue = [a for c in courses for a in c.get("assignments_overdue", [])]
    section("Overdue, still submittable", [
        _line(a, f" | closes {a['lock_local']}" if a.get("lock_local") else "") for a in overdue
    ])

    opens = [a for c in courses for a in c.get("assignments_open", [])]
    opens.sort(key=lambda a: a.get("due_at") or "9999")
    this_week = [a for a in opens if parse_dt(a.get("due_at")) and parse_dt(a["due_at"]) <= week_end]
    later = [a for a in opens if a not in this_week]
    section("Due in the next 7 days", [_line(a) for a in this_week])
    section("Due later", [_line(a) for a in later])

    undated_rows = []
    for c in courses:
        und = c.get("assignments_undated", [])
        if not und:
            continue
        available = [a for a in und if (a.get("module") or {}).get("module_state") != "locked"
                     and not (a.get("module") or {}).get("locked")]
        locked = len(und) - len(available)
        for a in available:
            mod = (a.get("module") or {}).get("module")
            undated_rows.append(_line(a, f" | module: {mod}" if mod else ""))
        if locked:
            undated_rows.append(f"- **{c['name']}** | {locked} more in locked modules")
    section("Undated, paced by modules (available now, in module order)", undated_rows)

    closed_rows = []
    for c in courses:
        for a in c.get("assignments_missing_closed", []):
            why = a.get("lock_explanation") or (f"closed {a['lock_local']}" if a.get("lock_local")
                                                else "locked")
            closed_rows.append(_line(a, f" | {why}"))
    section("Missing and closed (only possible if the teacher reopens it)", closed_rows)

    resub_rows = []
    for c in courses:
        for a in c.get("assignments_completed", []):
            r = a.get("resubmit")
            if not r:
                continue
            note = r.get("latest_teacher_comment") or ""
            note = (note[:140] + "...") if len(note) > 140 else note
            note = note.replace("\n", " ")
            resub_rows.append(_line(a, f" | {score_text(a)} | {', '.join(r['reasons'])}"
                                       f" | attempts left: {r.get('attempts_left')}"
                                       + (f" | teacher: \"{note}\"" if note else "")))
    section("Resubmission opportunities (graded, still open, low score or feedback)", resub_rows)

    section("Recently graded", [
        _line(a, f" | {score_text(a)}") for c in courses for a in c.get("assignments_completed", [])
        if a.get("status") == "graded" and a.get("bucket") == "recent_completed"
    ])

    section("Submitted, waiting on a grade", [
        _line(a) for c in courses for a in c.get("assignments_completed", [])
        if a.get("status") == "submitted_awaiting_grade"
    ])

    section("Check with teacher (due date is from before this term)", [
        _line(a) for c in courses for a in c.get("assignments_old_dates", [])
    ])

    manual = []
    for c in courses:
        for key in ("assignments_overdue", "assignments_open", "assignments_undated",
                    "assignments_old_dates"):
            for a in c.get(key, []):
                if a.get("external_tool_note"):
                    manual.append(f"- {a['name']} (id {a['id']}): external tool, open and paste it")
                for f in a.get("attachments", []):
                    if f.get("note") and not f.get("text"):
                        manual.append(f"- {a['name']} (id {a['id']}): file "
                                      f"\"{f.get('name', f.get('file_id'))}\": {f['note']}")
                for g in a.get("google_docs", []):
                    if g.get("note") and not g.get("text"):
                        manual.append(f"- {a['name']} (id {a['id']}): Google link {g['url']}: "
                                      f"{g['note']}")
    section("Needs a manual paste or upload", manual)

    outside_rows = []
    for c in courses:
        links = c.get("outside_canvas_links", [])
        for l in links[:MAX_OUTSIDE_LINKS_IN_SUMMARY]:
            outside_rows.append(f"- **{c['name']}** | {l.get('module', '')} | "
                                f"{l.get('title')} | {l['url']}")
        if len(links) > MAX_OUTSIDE_LINKS_IN_SUMMARY:
            outside_rows.append(f"- **{c['name']}** | {len(links) - MAX_OUTSIDE_LINKS_IN_SUMMARY} "
                                f"more in {OUTPUT_FILE}")
    section("Links to work outside Canvas (from modules)", outside_rows)

    course_rows = []
    for c in courses:
        active = any(c.get(k) for k in ("assignments_overdue", "assignments_open",
                                        "assignments_undated", "assignments_completed",
                                        "assignments_missing_closed", "outside_canvas_links"))
        course_rows.append(f"- {c['name']} | id {c['id']} | {c.get('term', 'no term')}"
                           + ("" if active else " | no current work"))
    course_rows.append("- If two of these are the same class, add the unused one's id to "
                       "EXCLUDE_COURSES in .env")
    section("Courses pulled", course_rows)

    if dump.get("errors"):
        L.append(f"_{len(dump['errors'])} request errors, see the `errors` list in {OUTPUT_FILE}._")
    return "\n".join(L).rstrip() + "\n"


def summary_mode(args):
    t0 = time.time()
    load_manifest()
    now = datetime.now(timezone.utc)
    cutoff = now - timedelta(days=args.days)

    courses = list_courses()
    enrollments = list_student_enrollments()
    if enrollments:
        courses = [c for c in courses if c["id"] in enrollments]
    courses = [c for c in courses if c["id"] not in EXCLUDE_COURSES]
    if args.course:
        courses = [c for c in courses if c["id"] == args.course]

    print(f"Lookback: {args.days} days. Courses:")
    for c in courses:
        print(f"  {c['id']}  {c.get('name')}  ({(c.get('term') or {}).get('name')})")
    print()

    results = []
    with ThreadPoolExecutor(ITEM_WORKERS) as item_pool, \
            ThreadPoolExecutor(COURSE_WORKERS) as course_pool:
        futures = [(c, course_pool.submit(build_course, c, enrollments.get(c["id"]),
                                          now, cutoff, item_pool)) for c in courses]
        for c, f in futures:
            try:
                results.append(f.result())
            except Exception as e:
                print(f"  !! course {c.get('id')} failed: {e}")
                _log_error(f"course {c.get('id')}", e)
    results.sort(key=lambda c: c.get("name") or "")

    dump = {
        "generated_at": now.isoformat(),
        "generated_local": fmt_local(now),
        "canvas_url": CANVAS_URL,
        "lookback_days": args.days,
        "courses": results,
        "errors": ERRORS,
    }
    write_json(args.out, dump)
    write_text(SUMMARY_FILE, build_summary(dump, now))
    save_manifest()

    print(f"\nWrote {SUMMARY_FILE}, {args.out}, and {ATTACH_DIR}/")
    print(f"  {STATS['requests']} requests, {STATS['files_downloaded']} files downloaded, "
          f"{STATS['files_reused']} reused, {STATS['gdocs_fetched']} Google Docs, "
          f"{len(ERRORS)} errors, {time.time() - t0:.1f}s")


def full_mode(args):
    now = datetime.now(timezone.utc)
    aid = args.full
    load_manifest()
    print(f"Looking up assignment {aid}...")
    courses = [c for c in list_courses() if c["id"] not in EXCLUDE_COURSES]

    def probe(c):
        a = get_one(f"/api/v1/courses/{c['id']}/assignments/{aid}", [("include[]", "submission")])
        return (c, a) if a else None

    with ThreadPoolExecutor(ITEM_WORKERS) as pool:
        found = next((r for r in pool.map(probe, courses) if r), None)
    if not found:
        print("  not found in any active course")
        sys.exit(1)

    course, a = found
    cid = course["id"]
    print(f"  found in {course.get('name')}, pulling everything...")
    amap, qmap, dmap = module_maps(list_modules(cid))
    topic_id = (a.get("discussion_topic") or {}).get("id")
    minfo = amap.get(a["id"]) or qmap.get(a.get("quiz_id")) or dmap.get(topic_id)
    groups = as_list(get_all(f"/api/v1/courses/{cid}/assignment_groups"))
    sub = fetch_submission(cid, aid) or a.get("submission")
    built = build_assignment(cid, course.get("name"), a, sub, "full", now, minfo,
                             {g["id"]: g.get("name") for g in groups if g.get("id")})

    out = FULL_FILE_TMPL.format(aid=aid)
    write_json(out, {"assignment": built, "errors": ERRORS})
    save_manifest()
    print(f"Wrote {out} ({STATS['requests']} requests)")


def main():
    p = argparse.ArgumentParser(description="Pull Canvas data.")
    p.add_argument("--full", type=int, metavar="ASSIGNMENT_ID", help="deep-dive one assignment")
    p.add_argument("--days", type=int, default=DEFAULT_LOOKBACK_DAYS,
                   help=f"window for recently completed work (default {DEFAULT_LOOKBACK_DAYS})")
    p.add_argument("--course", type=int, metavar="COURSE_ID", help="only pull this course")
    p.add_argument("--raw", action="store_true", help="also keep raw Canvas assignment objects")
    p.add_argument("--out", default=OUTPUT_FILE, help="JSON output file name")
    args = p.parse_args()

    CFG["raw"] = args.raw
    if args.full:
        full_mode(args)
    else:
        summary_mode(args)


if __name__ == "__main__":
    main()
